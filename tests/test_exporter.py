import io

from openpyxl import load_workbook

from paper_recycle_agent import exporter, regions


def _load(code="CN", year=2024):
    data = exporter.build_workbook(code, year)
    assert data[:2] == b"PK"  # xlsx 是 zip 容器
    return load_workbook(io.BytesIO(data))


def test_sheets_present():
    wb = _load()
    assert wb.sheetnames == ["省级统计", "分类回收", "历年回收", "废纸进口", "数据说明"]


def test_province_sheet_rows():
    ws = _load()["省级统计"]
    # 标题 + 备注 + 表头 = 3 行，之后 34 省 + 合计
    assert ws.max_row == 3 + 34 + 1
    values = [ws.cell(row=r, column=2).value for r in range(4, ws.max_row + 1)]
    assert "山东省" in values
    assert ws.cell(row=ws.max_row, column=1).value == "合计"


def test_province_sheet_total():
    ws = _load()["省级统计"]
    total_row = ws.max_row
    assert ws.cell(row=total_row, column=1).value == "合计"
    assert round(ws.cell(row=total_row, column=4).value, 1) == 6971.0


def test_category_sheet_eight_rows():
    ws = _load()["分类回收"]
    names = [ws.cell(row=r, column=2).value for r in range(4, 12)]
    assert len(names) == 8
    assert "废瓦楞纸箱类" in names
    assert ws.cell(row=12, column=1).value == "合计"


def test_history_sheet_years():
    ws = _load()["历年回收"]
    years = [ws.cell(row=r, column=1).value for r in range(4, ws.max_row + 1)]
    assert 2014 in years and 2025 in years
    assert years == sorted(years)


def test_import_sheet():
    ws = _load()["废纸进口"]
    pairs = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value
             for r in range(4, ws.max_row + 1)}
    assert pairs[2017] == 2572.0
    assert pairs[2024] == 0.0


def test_province_scoped_export():
    wb = _load("370000")
    ws = wb["省级统计"]
    # 只有 1 行数据
    assert ws.max_row == 3 + 1
    assert ws.cell(row=4, column=2).value == "山东省"


def test_notes_sheet_mentions_scope():
    ws = _load()["数据说明"]
    text = "\n".join(str(ws.cell(row=r, column=1).value or "") for r in range(1, ws.max_row + 1))
    assert "市级、县级" in text
    assert "6971" in text


def test_filename():
    assert exporter.filename("CN", 2024) == "废纸回收统计_全国_2024.xlsx"
    assert "山东省" in exporter.filename("370000", 2024)


def test_export_unknown_code_falls_back_to_national():
    wb = _load("NOT_A_CODE")
    assert wb["省级统计"].max_row == 3 + 34 + 1
